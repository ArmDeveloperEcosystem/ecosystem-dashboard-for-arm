"""Word, JSON, and CSV reporting with linked, scoped authoritative evidence."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .text import display_text, json_dumps

STATUS = {
    "supported": "Arm64 supported",
    "gap": "Arm64 support gap",
    "unknown": "Arm64 support unclear",
}


def clean(value):
    return display_text(value)


def hyperlink(paragraph, label, url):
    """Write evidence as a link only when its destination is unambiguous HTTPS."""
    url = str(url)
    if clean(url) != url:
        # A display escape is not a replacement URL. Keep the original destination
        # in raw JSON; show its escaped form here without creating a relationship.
        paragraph.add_run(clean(label) + " (" + clean(url) + ")")
        return
    try:
        parts = urlsplit(url)
        valid = (
            parts.scheme == "https"
            and bool(parts.hostname)
            and not parts.username
            and not parts.password
            and not re.search(r"[\\\x00-\x20\x7f]", url)
        )
    except ValueError:
        valid = False
    if not valid:
        paragraph.add_run(clean(label))
        return
    relation = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "176A84")
    prop.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    prop.append(underline)
    run.append(prop)
    text = OxmlElement("w:t")
    text.text = clean(label)
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    pr = t._tbl.tblPr
    for tag, attrs in (
        ("tblW", {"w": 9360, "type": "dxa"}),
        ("tblInd", {"w": 120, "type": "dxa"}),
    ):
        elem = pr.find(qn("w:" + tag))
        if elem is None:
            elem = OxmlElement("w:" + tag)
            pr.append(elem)
        for key, value in attrs.items():
            elem.set(qn("w:" + key), str(value))
    margins = OxmlElement("w:tblCellMar")
    for key, value in {"top": 80, "bottom": 80, "start": 120, "end": 120}.items():
        child = OxmlElement("w:" + key)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")
        margins.append(child)
    pr.append(margins)
    for i, width in enumerate(widths):
        t._tbl.tblGrid.gridCol_lst[i].set(qn("w:w"), str(width))
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = clean(value)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F2F4F7")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = clean(value)
    for row in t.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), str(width))
            for p in cell.paragraphs:
                p.style = doc.styles["Table Text"]
    return t


def document_style(doc):
    """standard_business_brief + memo_masthead; named compact table/citation override."""
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = (
        section.right_margin
    ) = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    tokens = {
        "Normal": (11, "202B33", 0, 6, 1.10),
        "Title": (26, "123B4A", 0, 8, 1.10),
        "Subtitle": (12, "536873", 0, 10, 1.10),
        "Heading 1": (16, "2E74B5", 16, 8, 1.10),
        "Heading 2": (13, "2E74B5", 12, 6, 1.10),
        "Heading 3": (12, "1F4D78", 8, 4, 1.10),
        "Table Text": (9, "202B33", 0, 4, 1.10),
        "Evidence": (9, "536873", 4, 4, 1.10),
        "Metadata": (10, "536873", 0, 3, 1.10),
        "Table Summary": (10, "536873", 4, 3, 1.10),
    }
    for name, (size, color, before, after, line) in tokens.items():
        style = (
            doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
        )
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line
        style.paragraph_format.widow_control = True
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.style = "Metadata"
    header.text = "ARM LINUX ECOSYSTEM  |  INTERNAL OPPORTUNITY REVIEW"
    footer = section.footer.paragraphs[0]
    footer.style = "Metadata"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Internal research PoC  |  Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Linux Arm64 support gap discovery"
    doc.core_properties.subject = (
        "Evidence-backed internal enablement opportunity review"
    )
    doc.core_properties.author = "Arm Linux Ecosystem PoC"


def popularity_signals(finding):
    """Keep every metric's date and unit; old saved observations remain readable."""
    supplied = finding.get("popularity_signals")
    if isinstance(supplied, list) and supplied:
        return supplied
    metadata = finding.get("metadata", {})
    if finding.get("source") == "github":
        return [
            {
                "name": "GitHub stars",
                "value": metadata.get("stargazers_count"),
                "unit": "stars",
                "period": "cumulative snapshot",
                "observed_at": finding.get("checked_at", "Unavailable"),
                "source_url": "https://github.com/" + finding["name"],
            }
        ]
    return [
        {
            "name": "Docker Hub pulls",
            "value": metadata.get("pull_count"),
            "unit": "pulls",
            "period": "cumulative snapshot",
            "observed_at": finding.get("checked_at", "Unavailable"),
            "source_url": "https://hub.docker.com/r/" + finding["name"],
        }
    ]


def signal_text(signal):
    value = signal.get("value")
    if value is None:
        rendered = "unavailable"
    elif isinstance(value, int) and not isinstance(value, bool):
        rendered = f"{value:,}"
    else:
        rendered = clean(value)
    return (
        f"{signal.get('name', 'Popularity signal')}: {rendered}"
        f" ({signal.get('period') or 'period not supplied'})."
        f" Observed: {signal.get('observed_at') or 'unavailable'}."
    )


def csv_safe(value):
    """Protect formula-like source strings, including whitespace-prefixed formulas."""
    if not isinstance(value, str):
        return value
    value = clean(value)
    if value.startswith(("\t", "\r", "\n")) or re.match(r"^\s*[=+@-]", value):
        return "'" + value
    return value


def write_csv(path, summary):
    fields = [
        "record_type",
        "candidate_id",
        "name",
        "source",
        "status",
        "scope",
        "checked_at",
        "next_check_at",
        "selection_reason",
        "popularity_signals",
        "catalog_tracked",
        "investigated_before",
        "previous_status",
        "evidence_changed",
        "reason",
        "recommended_action",
        "evidence_urls",
        "coverage_review_required",
        "coverage_review_reasons",
        "coverage_inventory_complete",
        "coverage_evidence_urls",
        "assessment_coverage",
    ]
    with open(path, "w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for record_type, records in (
            ("checked_this_run", summary["findings"]),
            ("historical_not_rechecked", summary.get("retained_findings", [])),
        ):
            for finding in records:
                values = {key: finding.get(key, "") for key in fields}
                values["record_type"] = record_type
                values["popularity_signals"] = json.dumps(
                    popularity_signals(finding), ensure_ascii=False
                )
                values["evidence_urls"] = " | ".join(
                    e["url"] for e in finding.get("evidence", [])
                )
                coverage = finding.get("assessment_coverage")
                if coverage is not None:
                    values.update(
                        coverage_review_required=coverage.get("review_required"),
                        coverage_review_reasons=" | ".join(
                            coverage.get("review_reasons", [])
                        ),
                        coverage_inventory_complete=coverage.get("inventory_complete"),
                        coverage_evidence_urls=" | ".join(
                            coverage.get("evidence_urls", [])
                        ),
                        assessment_coverage=json_dumps(coverage),
                    )
                writer.writerow({k: csv_safe(v) for k, v in values.items()})


def catalog_label(value):
    return (
        "Already represented"
        if value is True
        else "No exact repository/image URL match"
        if value is False
        else "Comparison unavailable"
    )


def add_coverage(doc, finding):
    """Expose inventory questions without reclassifying the evidence verdict."""
    coverage = finding.get("assessment_coverage")
    if coverage is None:
        if finding.get("source") == "github":
            doc.add_paragraph(
                "Structured assessment coverage was not recorded for this finding. "
                "Refer to its dated scope and original evidence.",
                "Evidence",
            )
        return
    doc.add_heading("Assessment coverage", level=3)
    complete = coverage.get("inventory_complete")
    doc.add_paragraph(
        "Inventory completeness: "
        + (
            "complete for the inspected scope"
            if complete is True
            else "incomplete"
            if complete is False
            else "not established"
        )
        + ". Coverage review does not change the support verdict.",
        "Metadata",
    )
    if coverage.get("review_required"):
        doc.add_paragraph(
            "Coverage review needed: "
            + clean(" ".join(coverage.get("review_reasons", []))),
            "Metadata",
        )
    supported = coverage.get("supported_artifacts", [])
    doc.add_paragraph(
        "Verified advertised Linux Arm64 artifacts: " + clean("; ".join(supported))
        if supported
        else "No advertised Linux Arm64 artifact was established in this inventory.",
        "Evidence",
    )
    remaining = coverage.get("remaining_inventory", [])
    doc.add_paragraph(
        f"Remaining collected inventory: {len(remaining)} records. "
        "Other architectures or component names alone do not establish a missing Arm64 component.",
        "Evidence",
    )
    if remaining:
        table(
            doc,
            ["Collected asset", "Inventory assessment"],
            [
                [
                    item.get("name") or "(unnamed asset)",
                    str(item.get("assessment", "unassessed")).replace("_", " ")
                    + (": " + str(item["reason"]) if item.get("reason") else ""),
                ]
                for item in remaining
            ],
            [5100, 4260],
        )
    for url in coverage.get("evidence_urls", []):
        hyperlink(doc.add_paragraph(style="Evidence"), "Inventory evidence", url)


def add_finding(doc, finding):
    """A readable, scoped decision record; it never widens a collector verdict."""
    f = finding
    first_paragraph = len(doc.paragraphs)
    doc.add_heading(clean(f["name"]), level=2)
    doc.add_paragraph("Checked scope: " + clean(f["scope"]), "Metadata")
    doc.add_paragraph(
        clean(
            f"Evidence checked: {f['checked_at']} | Next refresh: {f['next_check_at']}"
        ),
        "Metadata",
    )
    doc.add_paragraph(
        "Selection: "
        + clean(f.get("selection_reason") or "Retained investigation queue"),
        "Metadata",
    )
    for signal in popularity_signals(f):
        p = doc.add_paragraph(style="Metadata")
        hyperlink(p, signal_text(signal), signal.get("source_url", ""))
    metadata = f.get("metadata", {})
    maintenance = []
    if metadata.get("release_published_at"):
        maintenance.append(
            "Release published: " + str(metadata["release_published_at"])
        )
    if metadata.get("archived") is not None:
        maintenance.append(
            "Repository archived: " + ("yes" if metadata["archived"] else "no")
        )
    if maintenance:
        doc.add_paragraph(clean(" | ".join(maintenance)), "Metadata")
    if metadata.get("publisher_recognition"):
        doc.add_paragraph(
            "Publisher context: "
            + clean(metadata["publisher_recognition"])
            + ". This is not Arm64 certification.",
            "Metadata",
        )
    doc.add_paragraph(
        "Dashboard: "
        + catalog_label(f.get("catalog_tracked"))
        + ". Investigation: "
        + ("refresh of saved work" if f.get("investigated_before") else "first check")
        + ".",
        "Metadata",
    )
    doc.add_paragraph(clean(f["reason"]))
    p = doc.add_paragraph()
    p.add_run("Recommended action: ").bold = True
    p.add_run(clean(f["recommended_action"]))
    cited = [
        e
        for e in f.get("evidence", [])
        if e["kind"]
        in {
            "release_asset_inventory",
            "release_artifact",
            "dockerhub_tag_platforms",
            "oci_manifest",
            "oci_image_config",
        }
    ]
    if not cited:
        cited = f.get("evidence", [])[:2]
    for e in cited[:4]:
        p = doc.add_paragraph(style="Evidence")
        hyperlink(p, e["kind"].replace("_", " ").capitalize(), e["url"])
        excerpt = e["excerpt"]
        if e["kind"] == "release_asset_inventory" and len(e.get("asset_names", [])) > 6:
            completeness = "yes" if e.get("complete") is True else "not established"
            excerpt = (
                f"Complete inventory: {completeness}. Published assets: {len(e['asset_names'])}. "
                "Full filename inventory is retained in the linked source and JSON."
            )
        if e["kind"] == "oci_manifest" and e.get("manifest", {}).get("manifests"):
            from .evidence import manifest_platforms

            platforms, _ = manifest_platforms(e["manifest"])
            excerpt = "Runtime platforms: " + "; ".join(
                f"{v.get('os') or 'missing'}/{v.get('architecture') or 'missing'}"
                for v in platforms
                if not v.get("attestation") and not v.get("non_runtime_artifact")
            )
            excluded = sum(bool(v.get("attestation")) for v in platforms)
            if excluded:
                excerpt += f". Explicit attestation descriptors excluded: {excluded}."
            non_runtime = sum(bool(v.get("non_runtime_artifact")) for v in platforms)
            if non_runtime:
                excerpt += (
                    f". Non-runtime artifact descriptors excluded: {non_runtime}."
                )
        p.add_run(
            ": "
            + clean(excerpt[:450])
            + ("... [full evidence in JSON]" if len(excerpt) > 450 else "")
        )
    if not cited:
        doc.add_paragraph(
            "No usable evidence was collected. This is an unresolved check, not an unsupported verdict.",
            "Evidence",
        )
    if len(cited) > 4:
        doc.add_paragraph(
            f"{len(cited) - 4} additional evidence records retained in JSON.",
            "Evidence",
        )
    add_ai_review(doc, f)
    if f.get("failures"):
        doc.add_paragraph(
            "Collection issue: " + clean("; ".join(f["failures"])), "Evidence"
        )
    # Keep each bounded decision record together when it fits on one page.
    # Word can still paginate an unusually long AI note/evidence record.
    for paragraph in doc.paragraphs[first_paragraph:-1]:
        paragraph.paragraph_format.keep_with_next = True
    add_coverage(doc, f)


def add_ai_review(doc, f):
    review = f.get("ai_review", {})
    if review.get("status") == "completed":
        doc.add_paragraph(
            "AI advisory note (requires human review): " + clean(review["note"])
        )
        for citation in review.get("citations", []):
            p = doc.add_paragraph(style="Evidence")
            hyperlink(p, "Evidence-matched quotation", citation["url"])
            p.add_run(": " + clean(citation["quote"]))
    elif review.get("status") in {"pending", "failed_validation", "no_evidence"}:
        reason = clean(review.get("reason") or "").rstrip(".!?;: \t\r\n")
        doc.add_paragraph(
            "AI interpretation: "
            + review["status"].replace("_", " ")
            + (". " + reason if reason else "")
            + ". The metadata finding remains separate from AI interpretation.",
            "Evidence",
        )


def add_history(doc, summary):
    doc.add_heading("Saved findings: not rechecked this run", level=1)
    retained = summary.get("retained_findings", [])
    if not retained:
        doc.add_paragraph(
            "No historical findings are carried forward without rechecking in this report."
        )
        return
    doc.add_paragraph(
        f"{len(retained)} earlier findings are retained with their original observation dates. "
        "They are excluded from this run's counts and must not be treated as newly verified. "
        "Continue the existing investigation when follow-up is needed."
    )
    for f in sorted(
        retained,
        key=lambda v: (
            {"gap": 0, "unknown": 1, "supported": 2}[v["status"]],
            v["name"],
        ),
    ):
        first_paragraph = len(doc.paragraphs)
        doc.add_heading(clean(f["name"]) + " | " + STATUS[f["status"]], level=2)
        doc.add_paragraph("Historical scope: " + clean(f["scope"]), "Metadata")
        doc.add_paragraph(
            clean(
                f"Originally checked: {f['checked_at']} | Refresh due: {f['next_check_at']}"
            ),
            "Metadata",
        )
        doc.add_paragraph(
            "Selection: " + clean(f.get("selection_reason") or "Saved investigation"),
            "Metadata",
        )
        for signal in popularity_signals(f):
            p = doc.add_paragraph(style="Metadata")
            hyperlink(p, signal_text(signal), signal.get("source_url", ""))
        doc.add_paragraph(clean(f["reason"]))
        doc.add_paragraph("Next action: " + clean(f["recommended_action"]), "Metadata")
        evidence = sorted(
            f.get("evidence", []),
            key=lambda e: {
                "oci_manifest": 0,
                "release_asset_inventory": 1,
                "dockerhub_tag_platforms": 2,
            }.get(e["kind"], 9),
        )
        for e in evidence[:2]:
            p = doc.add_paragraph(style="Evidence")
            hyperlink(p, "Original " + e["kind"].replace("_", " "), e["url"])
        if not evidence:
            doc.add_paragraph(
                "Original check did not establish usable evidence.", "Evidence"
            )
        add_ai_review(doc, f)
        # Retain the historical scope and its final advisory disclosure together
        # when the bounded record fits on a page, as for a fresh finding above.
        for paragraph in doc.paragraphs[first_paragraph:-1]:
            paragraph.paragraph_format.keep_with_next = True
        add_coverage(doc, f)


def add_work_list(doc, title, items):
    doc.add_heading(title, level=2)
    if not items:
        doc.add_paragraph("None recorded for this run.", "Metadata")
        return
    for item in items[:25]:
        label = (
            item.get("candidate_id")
            or item.get("query")
            or item.get("source")
            or item.get("candidate")
            or "Run"
        )
        reason = (
            item.get("reason")
            or item.get("selection_reason")
            or "Awaiting investigation"
        )
        detail = f"{label}: {reason}"
        if item.get("count") is not None:
            detail += f" (records: {item['count']})"
        if item.get("next_check_at"):
            detail += ". Next eligible: " + item["next_check_at"]
        doc.add_paragraph(clean(detail), "Evidence")
    if len(items) > 25:
        doc.add_paragraph(
            f"Showing 25 of {len(items)} records. All {len(items)} records are available in the accompanying JSON.",
            "Evidence",
        )


def write_reports(summary):
    """Write one consistent report set; fresh counts never include historical rows."""
    paths = summary["report_paths"]
    Path(paths["json"]).write_text(json_dumps(summary, indent=2), encoding="utf-8")
    write_csv(paths["csv"], summary)
    doc = Document()
    document_style(doc)
    doc.add_paragraph("Linux Arm64\nsupport opportunities", "Title")
    doc.add_paragraph(
        "Internal evidence review | Selected releases and container tags", "Subtitle"
    )
    doc.add_paragraph(
        clean(
            "Run started: " + summary["generated_at"] + " | Run: " + summary["run_id"]
        ),
        "Metadata",
    )
    doc.add_heading("Decision summary", level=1)
    c = summary["counts"]
    doc.add_paragraph(
        f"This run checked {c['investigated']} release or container-tag scopes: "
        f"{c['gap']} with an identified Linux Arm64 distribution gap, "
        f"{c['unknown']} unclear, and {c['supported']} with advertised Linux Arm64 artifacts. "
        "Review scoped gaps for enablement priority; investigate unclear results before deciding whether a gap exists."
    )
    table(
        doc,
        ["Support gaps", "Unclear", "Supported", "Checked this run"],
        [[c["gap"], c["unknown"], c["supported"], c["investigated"]]],
        [2340] * 4,
    )
    doc.add_paragraph(
        f"First checks: {c['newly_investigated']}. Refreshed: {c['refreshed']}. "
        f"Earlier findings retained without rechecking: {len(summary.get('retained_findings', []))}. "
        f"Metadata requests: {c['requests']}. Collection/review issues: {c['failures']}.",
        "Table Summary",
    )
    coverage_current = sum(
        bool(f.get("assessment_coverage", {}).get("review_required"))
        for f in summary["findings"]
    )
    coverage_historical = sum(
        bool(f.get("assessment_coverage", {}).get("review_required"))
        for f in summary.get("retained_findings", [])
    )
    doc.add_paragraph(
        f"Coverage questions: {coverage_current} findings checked this run; "
        f"{coverage_historical} historical findings. These flags are separate from support status counts.",
        "Metadata",
    )
    scheduling = summary.get("scheduling")
    if scheduling:
        doc.add_paragraph(
            clean(
                f"Queue progress: {scheduling['status'].replace('_', ' ')}; "
                f"{scheduling['attempted']} attempted; {scheduling['deferred_due']} due scopes deferred. "
                + scheduling.get("reason", "")
            ),
            "Metadata",
        )
    if summary.get("queue") or c.get("pending_investigation"):
        doc.add_paragraph(
            f"Awaiting first investigation: {c.get('pending_investigation', len(summary.get('queue', [])))}. "
            "Remaining work stays in the saved queue for a later run.",
            "Metadata",
        )
    doc.add_paragraph(
        "Each result applies only to the named release artifacts or container tag. "
        "Unknown does not mean unsupported. These are metadata findings, not certification or runtime tests. "
        "Catalog membership is context and does not determine the verdict."
    )
    doc.add_paragraph(clean(summary["ai_review"]["disclosure"]), "Metadata")
    if summary["ai_review"].get("requested"):
        ai = summary["ai_review"]
        doc.add_paragraph(
            f"AI review this run: {ai.get('completed', 0)} completed of "
            f"{ai.get('eligible', 0)} findings with evidence; "
            f"{ai.get('failed', 0)} failed reviews. Historical findings retain their earlier AI status.",
            "Metadata",
        )
    doc.add_heading("What the findings mean", level=1)
    definitions = [
        (
            "Supported",
            "The inspected distribution advertises at least one Linux Arm64 artifact or runtime platform. Other components and runtime compatibility remain untested.",
        ),
        (
            "Identified gap",
            "A complete, interpretable inventory shows a missing Linux Arm64 distribution for the exact checked scope. It does not establish that the whole project cannot run on Arm.",
        ),
        (
            "Unclear / unknown",
            "Missing, incomplete or ambiguous evidence prevents a conclusion. This includes failed requests and bounded scans that cannot establish completeness.",
        ),
    ]
    for label, meaning in definitions:
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(meaning)
    doc.add_paragraph(
        "Human review is required before assigning follow-up or proposing catalog changes. "
        "No finding automatically changes the dashboard or contacts a maintainer.",
        "Metadata",
    )
    doc.add_heading("Checked this run", level=1)
    if not summary["findings"]:
        doc.add_paragraph(
            "No candidates were investigated in this run. Earlier findings remain historical; "
            "their statuses have not been newly verified or reclassified. See saved findings and queued work."
        )
    for status, title in (
        ("gap", "Identified gaps: confirm an enablement priority"),
        ("unknown", "Unclear support: resolve the evidence"),
        ("supported", "Supported distributions: retain the evidence"),
    ):
        findings = [f for f in summary["findings"] if f["status"] == status]
        if findings:
            doc.add_heading(f"{title} ({len(findings)})", level=1)
            for f in sorted(findings, key=lambda value: value["name"]):
                add_finding(doc, f)
    add_history(doc, summary)
    doc.add_heading("Coverage, selection and unfinished work", level=1)
    selection = summary.get("selection", {})
    queries = selection.get("github_queries", [])
    doc.add_paragraph(
        "This is a bounded sample of selected project repositories and container tags. "
        "It is not a census of open-source software or a global popularity ranking. "
        "Popularity helps select work; it is not proof of Arm64 support. Stars and pulls remain separate metrics. "
        "Unavailable statistics are not recorded as zero."
    )
    if selection:
        doc.add_paragraph(
            clean(
                f"Configured seeds: {len(selection.get('seeds', []))}. GitHub discovery queries: "
                + ("; ".join(queries) or "none")
                + f". Minimum search stars: {selection.get('minimum_stars', 'not specified')}."
            ),
            "Metadata",
        )
        for key, label in (
            ("repository_release_selection", "Release selection"),
            ("registry_selection", "Container selection"),
            ("priority_policy", "Queue priority"),
            ("retained_queue_policy", "Saved work"),
        ):
            if selection.get(key):
                doc.add_paragraph(label + ": " + clean(selection[key]), "Metadata")
    if selection.get("source_records_fetched") is not None:
        doc.add_paragraph(
            f"Discovery source records: {selection['source_records_fetched']} collected; "
            f"{selection.get('source_records_examined', 0)} examined; "
            f"{selection.get('source_candidates_selected', 0)} candidates selected.",
            "Metadata",
        )
    limits = summary["limits"]
    main_limits = [
        ("max_candidates", "Scopes investigated"),
        ("max_discovered", "New search candidates"),
        ("max_source_records", "Source records collected"),
        ("max_requests", "Metadata requests"),
        ("max_seconds", "Run seconds"),
        ("max_queries", "Search queries"),
    ]
    workload_table = table(
        doc,
        ["Workload limit", "Maximum per run"],
        [[label, limits[key]] for key, label in main_limits if key in limits],
        [6500, 2860],
    )
    # This small table is one workload summary; keep it on one page when it fits.
    for row in workload_table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph(
        "These are maximum workloads, not promised findings. A request limit is not a package count: "
        "one candidate may need several requests. Full request, response-size and pagination limits are retained in JSON.",
        "Evidence",
    )
    doc.add_paragraph(
        clean(
            "Refresh intervals: "
            + "; ".join(
                f"{STATUS.get(key, key)} every {value} hours"
                for key, value in summary["refresh_hours"].items()
            )
            + ". Previously checked candidates retain their original dates until refreshed."
        ),
        "Metadata",
    )
    add_work_list(doc, "Queued for first investigation", summary.get("queue", []))
    if summary.get("quarantined_candidates"):
        add_work_list(
            doc,
            "Saved identities requiring manual review",
            summary["quarantined_candidates"],
        )
    skipped = summary.get("skipped", [])
    not_due = [item for item in skipped if item.get("reason") == "Not due for refresh"]
    other_skips = [item for item in skipped if item not in not_due]
    add_work_list(doc, "Saved work not yet due for refresh", not_due)
    add_work_list(
        doc, "Skipped, excluded or deferred work: recorded reasons", other_skips
    )
    add_work_list(doc, "Collection and review issues", summary.get("failures", []))
    boundaries_start = len(doc.paragraphs)
    doc.add_heading("Boundaries and traceability", level=1)
    for limitation in summary["limitations"]:
        doc.add_paragraph(clean(limitation), "Metadata")
    doc.add_paragraph(
        "The Word report supports review and decisions. The CSV contains current and retained historical findings, "
        "explicitly labeled by record type for filtering. JSON preserves every collected evidence record, "
        "selection detail, limit, queued/skipped item and collection issue. Historical rows keep their original dates. "
        "Saved state allows subsequent runs to continue work without reopening duplicate investigations.",
        "Evidence",
    )
    doc.add_paragraph(
        "Characters that cannot appear in Word XML are shown as visible Unicode escapes "
        "in Word, CSV and diagnostics. Invalid link destinations are shown as plain text. "
        "The accompanying JSON preserves the exact original source and model text.",
        "Evidence",
    )
    for paragraph in doc.paragraphs[boundaries_start:-1]:
        paragraph.paragraph_format.keep_with_next = True
    doc.save(paths["docx"])
