"""Report contracts: no widened verdicts, lost dates or unsafe source output."""

import csv
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from poc.discovery.report import csv_safe, hyperlink, write_reports


@pytest.fixture
def report_summary(tmp_path):
    def finding(name, status, source="github", checked="2026-09-24T12:00:00+00:00"):
        return {
            "candidate_id": f"{source}:example/{name}",
            "name": "example/" + name,
            "source": source,
            "status": status,
            "scope": f"example/{name} release v2.0: selected Linux binary assets only",
            "checked_at": checked,
            "next_check_at": "2026-10-01T12:00:00+00:00",
            "selection_reason": "Configured database-category seed",
            "reason": f"Evidence establishes only the {status} finding for v2.0.",
            "catalog_tracked": True,
            "investigated_before": False,
            "recommended_action": "Human review of this exact distribution; confirm existing ownership.",
            "evidence": [
                {
                    "kind": "release_asset_inventory",
                    "url": f"https://api.github.com/repos/example/{name}/releases/42/assets",
                    "excerpt": "Complete=True; Linux distribution inventory checked.",
                }
            ],
            "metadata": {"stargazers_count": 0, "archived": False},
            "popularity_signals": [
                {
                    "name": "GitHub stars",
                    "value": 0,
                    "unit": "stars",
                    "period": "cumulative snapshot",
                    "observed_at": checked,
                    "source_url": f"https://github.com/example/{name}",
                }
            ],
            "ai_review": {"status": "not_configured"},
            "failures": [],
        }

    fresh = [
        finding("gap", "gap"),
        finding("supported", "supported"),
        finding("unclear", "unknown"),
    ]
    fresh[-1]["evidence"] = []
    fresh[-1]["popularity_signals"][0]["value"] = None
    historical = finding("old-gap", "gap", checked="2026-09-10T10:00:00+00:00")
    historical["historical"] = True
    return {
        "report_paths": {
            ext: str(tmp_path / ("opportunities." + ext))
            for ext in ("json", "csv", "docx")
        },
        "generated_at": "2026-09-24T12:00:00+00:00",
        "run_id": "run-for-report-test",
        "counts": {
            "investigated": 3,
            "gap": 1,
            "unknown": 1,
            "supported": 1,
            "newly_investigated": 3,
            "refreshed": 0,
            "catalog_tracked": 3,
            "saved_observations": 4,
            "requests": 12,
            "failures": 1,
            "pending_investigation": 1,
        },
        "findings": fresh,
        "retained_findings": [historical],
        "saved_history": [],
        "selection": {
            "seeds": [{"source": "github", "name": "example/gap"}],
            "github_queries": ["topic:database"],
            "minimum_stars": 500,
        },
        "queue": [
            {"candidate_id": "github:example/queued", "selection_reason": "Run limit"}
        ],
        "skipped": [
            {
                "candidate_id": "github:example/old-gap",
                "reason": "Not due for refresh",
                "next_check_at": "2026-10-01T12:00:00+00:00",
            },
            {
                "candidate_id": "github:example/queued",
                "reason": "Run candidate limit reached; saved queue retained",
            },
            {"source": "github", "reason": "HTTP 429; search deferred"},
        ],
        "failures": [{"source": "github", "reason": "HTTP 429; retry later"}],
        "limits": {
            "max_candidates": 8,
            "max_discovered": 2,
            "max_queries": 2,
            "max_requests": 50,
            "max_seconds": 180,
        },
        "refresh_hours": {"supported": 168, "gap": 168, "unknown": 24},
        "ai_review": {
            "disclosure": "No AI reviewer is configured; metadata verification only."
        },
        "limitations": [
            "No downloaded code was executed.",
            "Unknown results are not unsupported.",
        ],
        "state_path": "/private/local/reports/state.sqlite3",
    }


def report_text(path):
    doc = Document(path)
    return (
        "\n".join(p.text for p in doc.paragraphs)
        + "\n"
        + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    )


def test_exports_preserve_fresh_and_historical_verdict_scope_and_dates(report_summary):
    write_reports(report_summary)
    exported = json.loads(Path(report_summary["report_paths"]["json"]).read_text())
    assert exported == report_summary
    with open(report_summary["report_paths"]["csv"], newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 4
    assert {r["record_type"] for r in rows} == {
        "checked_this_run",
        "historical_not_rechecked",
    }
    old = next(r for r in rows if r["record_type"] == "historical_not_rechecked")
    assert old["status"] == "gap"
    assert old["checked_at"] == "2026-09-10T10:00:00+00:00"
    assert "v2.0" in old["scope"]
    assert json.loads(old["popularity_signals"])[0]["observed_at"] == old["checked_at"]
    fresh_counts = {
        status: sum(
            r["status"] == status and r["record_type"] == "checked_this_run"
            for r in rows
        )
        for status in ("gap", "unknown", "supported")
    }
    assert fresh_counts == {"gap": 1, "unknown": 1, "supported": 1}


def test_word_distinguishes_unknown_history_and_unfinished_work(report_summary):
    write_reports(report_summary)
    text = report_text(report_summary["report_paths"]["docx"])
    assert "This run checked 3 release or container-tag scopes: 1 with" in text
    assert "Unknown does not mean unsupported" in text
    assert "Unclear support: resolve the evidence (1)" in text
    assert "not an unsupported verdict" in text
    assert "Saved findings: not rechecked this run" in text
    assert "Originally checked: 2026-09-10T10:00:00+00:00" in text
    assert "Queued for first investigation" in text
    assert "Saved work not yet due for refresh" in text
    assert "HTTP 429; retry later" in text
    assert "one candidate may need several requests" in text
    assert report_summary["state_path"] not in text
    assert "No AI reviewer is configured" in text


def test_dated_signals_are_clickable_and_zero_is_not_missing(report_summary):
    write_reports(report_summary)
    path = report_summary["report_paths"]["docx"]
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
        relationships = archive.read("word/_rels/document.xml.rels").decode()
    assert "GitHub stars: 0 (cumulative snapshot)" in xml
    assert "GitHub stars: unavailable (cumulative snapshot)" in xml
    assert "Observed: 2026-09-10T10:00:00+00:00" in xml
    assert "https://github.com/example/old-gap" in relationships
    assert (
        "https://api.github.com/repos/example/gap/releases/42/assets" in relationships
    )
    assert "w:hyperlink" in xml
    doc = Document(path)
    for table in doc.tables:
        assert table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert sum(int(col.get(qn("w:w"))) for col in table._tbl.tblGrid) == 9360
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        for row in table.rows:
            assert (
                sum(
                    int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w")))
                    for cell in row.cells
                )
                == 9360
            )


def test_repeat_run_reports_no_new_findings_without_losing_history(report_summary):
    report_summary["retained_findings"] += report_summary["findings"]
    report_summary["findings"] = []
    for key in (
        "investigated",
        "gap",
        "unknown",
        "supported",
        "newly_investigated",
        "refreshed",
    ):
        report_summary["counts"][key] = 0
    write_reports(report_summary)
    text = report_text(report_summary["report_paths"]["docx"])
    assert "No candidates were investigated in this run" in text
    assert "4 earlier findings are retained" in text
    with open(report_summary["report_paths"]["csv"], newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 4
    assert all(row["record_type"] == "historical_not_rechecked" for row in rows)


@pytest.mark.parametrize(
    "source",
    ["=HYPERLINK(1)", "+SUM(1)", "-1+2", "@SUM(1)", "  =1", "\t=1", "\r=1", "\n=1"],
)
def test_source_strings_cannot_become_spreadsheet_formulas(source):
    assert csv_safe(source).startswith("'")
    assert csv_safe("Ordinary text") == "Ordinary text"
    assert csv_safe(0) == 0


@pytest.mark.parametrize(
    "url",
    [
        "http://github.com/example/repo",
        "javascript:alert(1)",
        "https://evil.example\\@github.com/repo",
        "https://user:password@github.com/repo",
        "https://github.com/\nexample",
        "https://[broken/",
    ],
)
def test_unsafe_hyperlinks_render_as_plain_labels(url):
    doc = Document()
    paragraph = doc.add_paragraph()
    hyperlink(paragraph, "Evidence", url)
    assert paragraph.text == "Evidence"
    assert not paragraph._p.findall(qn("w:hyperlink"))


def test_report_platform_summary_does_not_treat_sbom_as_runtime(report_summary):
    finding = report_summary["findings"][0]
    finding["evidence"] = [
        {
            "kind": "oci_manifest",
            "url": "https://registry-1.docker.io/v2/example/tool/manifests/v2",
            "excerpt": "Collected index",
            "manifest": {
                "schemaVersion": 2,
                "manifests": [
                    {"platform": {"os": "linux", "architecture": "amd64"}},
                    {
                        "artifactType": "application/spdx+json",
                        "platform": {"os": "linux", "architecture": "arm64"},
                    },
                    {
                        "annotations": {
                            "vnd.docker.reference.type": "attestation-manifest"
                        }
                    },
                ],
            },
        }
    ]
    finding["metadata"]["publisher_recognition"] = "Docker Official Image"
    write_reports(report_summary)
    with ZipFile(report_summary["report_paths"]["docx"]) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "Runtime platforms: linux/amd64." in xml
    assert "Runtime platforms: linux/amd64; linux/arm64" not in xml
    assert "Non-runtime artifact descriptors excluded: 1" in xml
    assert "Explicit attestation descriptors excluded: 1" in xml
    assert (
        "Publisher context: Docker Official Image. This is not Arm64 certification."
        in xml
    )


def test_coverage_questions_preserve_status_counts_and_match_all_exports(
    report_summary,
):
    coverage = {
        "kind": "github_release_assets",
        "inventory_complete": False,
        "supported_artifacts": ["server-linux-arm64.tgz"],
        "remaining_inventory": [
            {
                "name": "client-linux-amd64.tgz",
                "assessment": "other_linux_binary",
                "reason": "Other Linux architecture; no component inference",
            }
        ],
        "review_required": True,
        "review_reasons": ["Inventory pagination is incomplete."],
        "evidence_urls": [
            "https://api.github.com/repos/example/supported/releases/1/assets"
        ],
        "limitations": ["Runtime compatibility unassessed"],
    }
    report_summary["findings"][1]["assessment_coverage"] = coverage
    report_summary["retained_findings"][0]["assessment_coverage"] = coverage
    expected_counts = dict(report_summary["counts"])
    write_reports(report_summary)
    exported = json.loads(Path(report_summary["report_paths"]["json"]).read_text())
    assert exported["counts"] == expected_counts
    assert exported["findings"][1]["status"] == "supported"
    with open(report_summary["report_paths"]["csv"], newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 4
    supported = next(row for row in rows if row["status"] == "supported")
    assert supported["coverage_review_required"] == "True"
    assert json.loads(supported["assessment_coverage"]) == coverage
    assert supported["coverage_review_reasons"] == coverage["review_reasons"][0]
    assert supported["coverage_evidence_urls"] == coverage["evidence_urls"][0]
    assert rows[-1]["record_type"] == "historical_not_rechecked"
    text = report_text(report_summary["report_paths"]["docx"])
    assert (
        "Coverage questions: 1 findings checked this run; 1 historical findings" in text
    )
    assert "client-linux-amd64.tgz" in text
    assert "Coverage review needed: Inventory pagination is incomplete." in text
    assert "Historical scope:" in text
    assert report_summary["counts"] == expected_counts


def test_coverage_text_and_links_survive_invalid_unicode(report_summary):
    report_summary["findings"][1]["assessment_coverage"] = {
        "inventory_complete": True,
        "supported_artifacts": ["linux-arm64-\ud800.tgz"],
        "remaining_inventory": [
            {
                "name": "client\uffff",
                "assessment": "ambiguous",
                "reason": "unsafe\ud800",
            }
        ],
        "review_required": True,
        "review_reasons": ["ambiguous\ufffe"],
        "evidence_urls": ["https://github.com/example/tool/\uffff"],
    }
    write_reports(report_summary)
    text = report_text(report_summary["report_paths"]["docx"])
    assert "client\\uffff" in text
    assert "ambiguous\\ufffe" in text
    doc = Document(report_summary["report_paths"]["docx"])
    assert not any(
        "\\uffff" in str(rel.target_ref) or "\uffff" in str(rel.target_ref)
        for rel in doc.part.rels.values()
    )
    exported = json.loads(Path(report_summary["report_paths"]["json"]).read_text())
    assert (
        exported["findings"][1]["assessment_coverage"]
        == report_summary["findings"][1]["assessment_coverage"]
    )


def test_run_health_and_collection_counters_are_visible(report_summary):
    report_summary["scheduling"] = {
        "status": "no_progress",
        "attempted": 0,
        "deferred_due": 1,
        "reason": "Source budget exhausted; saved work retained.",
    }
    report_summary["selection"].update(
        source_records_fetched=100,
        source_records_examined=80,
        source_candidates_selected=2,
    )
    write_reports(report_summary)
    text = report_text(report_summary["report_paths"]["docx"])
    assert "Queue progress: no progress; 0 attempted; 1 due scopes deferred." in text
    assert "100 collected; 80 examined; 2 candidates selected" in text
