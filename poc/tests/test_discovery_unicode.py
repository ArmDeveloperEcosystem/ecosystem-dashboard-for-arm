"""Unusual public evidence must survive storage, history and all report formats."""

import csv
import json
import sqlite3
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from poc.discovery.pipeline import run_pipeline
from poc.discovery.report import hyperlink, write_reports
from poc.discovery.text import display_text, json_dumps

NOW = datetime(2026, 9, 24, tzinfo=timezone.utc)


class PublicGithub:
    def __init__(self, *, tag="v1", description="Public project", asset=None):
        self.requests_used = 0
        self.tag = tag
        self.description = description
        self.asset = asset or "app-linux-amd64.tar.gz"

    def get(self, url, **kwargs):
        self.requests_used += 1
        if url.endswith("/readme"):
            payload = {"encoding": "base64", "content": "", "sha": "a" * 40}
        elif url.endswith("/releases/latest"):
            payload = {"id": 1, "tag_name": self.tag, "body": "Linux artifacts are listed."}
        elif url.endswith("/assets"):
            payload = [{"name": self.asset, "size": 10, "state": "uploaded"}]
        else:
            payload = {
                "id": 2,
                "full_name": "public/project",
                "private": False,
                "visibility": "public",
                "description": self.description,
            }
        # JSON source responses can legally contain escaped lone surrogates.
        return json.loads(json.dumps(payload)), {}


def config(tmp_path):
    path = tmp_path / "config.yaml"
    path.write_text("seeds:\n - source: github\n   name: public/project\n")
    return path


def exported_summary(summary):
    return json.loads(Path(summary["report_paths"]["json"]).read_text("utf-8"))


def word_xml(summary):
    with ZipFile(summary["report_paths"]["docx"]) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def test_xml_boundaries_escape_visibly_and_valid_unicode_is_unchanged(tmp_path):
    # Separate lone surrogate code points: JSON decoders combine adjacent high /
    # low surrogate pairs into a valid supplementary character at ingestion.
    invalid = "\x00\x08\x0b\x0c\x1f\ud800|\udfff\ufffe\uffff"
    valid = "Café 中文 😀\t\n\r\u0020\ud7ff\ue000\ufffd\U00010000\U0010ffff"
    shown = display_text(invalid)
    assert shown == r"\u0000\u0008\u000b\u000c\u001f\ud800|\udfff\ufffe\uffff"
    assert display_text(valid) == valid
    assert json.loads(json_dumps({"raw": invalid + valid}))["raw"] == invalid + valid
    doc = Document()
    doc.add_paragraph(shown + valid)
    doc.save(tmp_path / "boundaries.docx")
    assert (tmp_path / "boundaries.docx").is_file()


@pytest.mark.parametrize("character", ["\uffff", "\ud800"])
def test_invalid_hyperlink_destination_is_visible_but_never_rewritten(
    tmp_path, character
):
    original = "https://github.com/public/project/releases/tag/v1" + character
    doc = Document()
    paragraph = doc.add_paragraph()
    hyperlink(paragraph, "Evidence " + character, original)
    assert paragraph.text == display_text(
        "Evidence " + character + " (" + original + ")"
    )
    assert not paragraph._p.findall(qn("w:hyperlink"))
    path = tmp_path / "invalid-link.docx"
    doc.save(path)
    with ZipFile(path) as archive:
        relationships = archive.read("word/_rels/document.xml.rels").decode()
    assert "github.com" not in relationships


def test_noncharacter_release_survives_fresh_and_historical_runs(tmp_path):
    path = config(tmp_path)
    output = tmp_path / "state"
    first = run_pipeline(path, output, now=NOW, http=PublicGithub(tag="v1\uffff"))
    finding = first["findings"][0]
    assert finding["candidate_id"] == "github:public/project"
    assert finding["status"] == "gap"
    assert "v1\uffff" in finding["scope"]
    assert exported_summary(first) == first
    assert r"v1\uffff" in word_xml(first)
    second_http = PublicGithub()
    second = run_pipeline(path, output, now=NOW + timedelta(hours=1), http=second_http)
    assert second_http.requests_used == 0
    assert second["counts"]["investigated"] == 0
    assert second["counts"]["saved_observations"] == 1
    assert second["retained_findings"][0]["scope"] == finding["scope"]
    assert exported_summary(second) == second
    assert r"v1\uffff" in word_xml(second)
    with open(second["report_paths"]["csv"], encoding="utf-8", newline="") as stream:
        row = next(csv.DictReader(stream))
    assert row["record_type"] == "historical_not_rechecked"
    assert row["scope"] == display_text(finding["scope"])


def test_legacy_failed_report_history_recovers_without_deleting_evidence(
    tmp_path, monkeypatch
):
    path = config(tmp_path)
    output = tmp_path / "state"

    def old_report_failure(summary):
        raise ValueError("All strings must be XML compatible")

    with monkeypatch.context() as patch:
        patch.setattr("poc.discovery.report.write_reports", old_report_failure)
        with pytest.raises(ValueError, match="XML compatible"):
            run_pipeline(path, output, now=NOW, http=PublicGithub(tag="v1\uffff"))

    state = output / "discovery.sqlite3"
    with sqlite3.connect(state) as db:
        original = json.loads(
            db.execute("SELECT result FROM observations").fetchone()[0]
        )
        # Recreate the previous version's successfully stored raw noncharacter
        # scope and non-ASCII JSON: upgrades must read this existing state too.
        legacy_json = json.dumps(original, ensure_ascii=False)
        db.execute(
            "UPDATE observations SET scope=?,result=?", (original["scope"], legacy_json)
        )
        db.execute("UPDATE candidates SET last_result=?", (legacy_json,))
        db.commit()
        before = db.execute(
            "SELECT candidate_id,scope,result FROM observations"
        ).fetchall()

    http = PublicGithub()
    recovered = run_pipeline(path, output, now=NOW + timedelta(hours=1), http=http)
    assert http.requests_used == 0
    assert recovered["counts"]["saved_observations"] == 1
    assert recovered["retained_findings"][0]["scope"] == original["scope"]
    assert r"v1\uffff" in word_xml(recovered)
    assert exported_summary(recovered) == recovered
    with sqlite3.connect(state) as db:
        assert (
            db.execute("SELECT candidate_id,scope,result FROM observations").fetchall()
            == before
        )
        assert db.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 2
    assert (
        json.loads((output / "latest.json").read_text("utf-8"))["run_id"]
        == recovered["run_id"]
    )


def test_surrogate_metadata_and_supported_scope_roundtrip_through_sqlite_and_exports(
    tmp_path,
):
    source = PublicGithub(
        description="Café 中文 😀 raw \ud800 metadata",
        asset="app-linux-arm64-\udfff.tar.gz",
    )
    result = run_pipeline(config(tmp_path), tmp_path / "state", now=NOW, http=source)
    finding = result["findings"][0]
    assert finding["candidate_id"] == "github:public/project"
    assert finding["status"] == "supported"
    assert finding["metadata"]["description"] == source.description
    assert source.asset in finding["scope"]
    assert exported_summary(result) == result
    with sqlite3.connect(result["state_path"]) as db:
        scope, raw = db.execute("SELECT scope,result FROM observations").fetchone()
        assert scope == display_text(finding["scope"])
        assert json.loads(raw) == finding
        assert (
            json.loads(db.execute("SELECT last_result FROM candidates").fetchone()[0])
            == finding
        )
    assert r"app-linux-arm64-\udfff.tar.gz" in word_xml(result)
    assert json.loads((tmp_path / "state/latest.json").read_text("utf-8")) == result


def test_surrogate_tag_uses_release_id_evidence_without_changing_raw_scope(tmp_path):
    result = run_pipeline(
        config(tmp_path), tmp_path / "state", now=NOW, http=PublicGithub(tag="v1\ud800")
    )
    finding = result["findings"][0]
    assert finding["status"] == "gap"
    assert finding["metadata"]["release_tag"] == "v1\ud800"
    assert "v1\ud800" in finding["scope"]
    release = next(e for e in finding["evidence"] if e["kind"] == "release_notes")
    assert release["url"] == "https://api.github.com/repos/public/project/releases/1"
    assert exported_summary(result) == result
    assert r"v1\ud800" in word_xml(result)


def test_model_advisory_surrogates_persist_without_changing_verdict(tmp_path):
    note = "Human review of this distribution. Café 中文 😀 \ud800 \uffff"
    source = PublicGithub(asset="app-linux-arm64-\udfff.tar.gz")

    def reviewer(evidence):
        item = next(e for e in evidence["evidence"] if e["kind"] == "release_artifact")
        return {
            "note": note,
            "citations": [{"url": item["url"], "quote": item["excerpt"]}],
        }

    result = run_pipeline(
        config(tmp_path), tmp_path / "state", now=NOW, http=source, reviewer=reviewer
    )
    finding = result["findings"][0]
    assert finding["status"] == "supported"
    assert finding["ai_review"]["status"] == "completed"
    assert finding["ai_review"]["note"] == note
    assert exported_summary(result) == result
    with sqlite3.connect(result["state_path"]) as db:
        stored = json.loads(db.execute("SELECT result FROM observations").fetchone()[0])
    assert stored == finding
    assert display_text(note) in word_xml(result)


def test_source_model_and_history_text_have_complete_presentation_boundaries(tmp_path):
    result = run_pipeline(
        config(tmp_path), tmp_path / "state", now=NOW, http=PublicGithub()
    )
    malformed = "Café 中文 😀 \ud800 \ufffe \uffff \x00"
    original = deepcopy(result)
    finding = result["findings"][0]
    finding["metadata"]["release_published_at"] = malformed
    finding["metadata"]["publisher_recognition"] = malformed
    finding["checked_at"] = malformed
    finding["next_check_at"] = malformed
    finding["selection_reason"] = malformed
    finding["reason"] = malformed
    finding["recommended_action"] = malformed
    finding["popularity_signals"][0].update(name=malformed, observed_at=malformed)
    finding["evidence"][2].update(
        url="https://github.com/public/project/\ud800", excerpt=malformed
    )
    finding["ai_review"] = {
        "status": "completed",
        "note": malformed,
        "citations": [
            {"url": "https://github.com/public/project/\uffff", "quote": malformed}
        ],
    }
    finding["failures"] = [malformed]
    result["retained_findings"] = [{**deepcopy(finding), "historical": True}]
    result["selection"]["github_queries"] = [malformed]
    result["queue"] = [{"candidate_id": finding["candidate_id"], "reason": malformed}]
    result["ai_review"]["disclosure"] = malformed
    result["limitations"].append(malformed)
    before = deepcopy(result)
    write_reports(result)
    assert result == before
    assert result["counts"] == original["counts"]
    assert exported_summary(result) == result
    xml = word_xml(result)
    assert "Café 中文 😀" in xml
    assert display_text(malformed) in xml
    with ZipFile(result["report_paths"]["docx"]) as archive:
        links = archive.read("word/_rels/document.xml.rels").decode()
    assert "\\ud800" not in links and "\\uffff" not in links
    with open(result["report_paths"]["csv"], encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 2
    assert all(row["reason"] == display_text(malformed) for row in rows)
